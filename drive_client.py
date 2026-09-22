import io
import os
import re

from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

SCOPES = [
    "https://www.googleapis.com/auth/drive.readonly",
    "https://www.googleapis.com/auth/spreadsheets.readonly",
]

GENERATION_PATTERN = re.compile(r"(\d{1,2})\s*기")

# Top-level team folders under the "Growth Hackers" Drive root, discovered by
# listing parentId='<root folder id>'. Folder sharing cascades to children,
# so only the root needs to be shared with the service account.
TEAM_FOLDERS = {
    "관리": "1L-U8I_INx4eL7ZQy8oVk_Bw9SpOnI_WN",
    "공통": "1DbtOQ5-MCAaqjzdBA5t189rCAKn6oAcC",
    "회장단": "1ON5pA7kNplNrI3i2I5CoBsk9wuxbBdD6",
    "NUT": "1Tu15zig-d67GN5hp0oeBBsqFU3IgkxY4",
    "DH": "1x-K9HuRkydqq7mry30W154y88kYQS9gK",
    "HR": "1A1MOzi579CyY4e0rWGEQyU7dpPubOwxA",
    "PR": "17jfwkRvFmpwHJfM13pJOvqMvc2qHbnZ9",
    "EDU": "1aefY780gVmKhRomyk62aFa9xOL_wz3Kp",
    "과거자료": "1bXYjeNUyayfW8bSyD66pz_5Yx3IlMtJP",
}

# Native Google formats need `export`, not `get_media` - they have no fixed byte content.
EXPORT_TEXT_MIME_TYPES = {
    "application/vnd.google-apps.document": "text/plain",
    "application/vnd.google-apps.presentation": "text/plain",
}

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

# The service account can see ~33,000 items (photos, zips, csvs, colab
# notebooks, JSON exports, ...) but only these types carry readable text worth
# embedding - about 2,500 of the 33,000. Filtering here, before any download,
# is what keeps indexing tractable.
SUPPORTED_MIME_TYPES = {
    *EXPORT_TEXT_MIME_TYPES.keys(),
    "application/pdf",
    DOCX_MIME,
    PPTX_MIME,
}

FOLDER_MIME = "application/vnd.google-apps.folder"

# Folders whose name matches these (case-insensitive substring) hold superseded
# material (e.g. "회장단-old" was found sitting inside the current 회장단 folder,
# and Drive has ~10 old dated copies of 정관 floating around). Excluded from
# recursive indexing so stale content doesn't dilute semantic_search, but
# search() still finds them if someone searches by name directly.
STALE_FOLDER_MARKERS = ["-old", "old버전", "구버전", "보관", "이전"]


def _is_stale_folder(name: str) -> bool:
    lowered = name.lower()
    return any(marker in lowered for marker in STALE_FOLDER_MARKERS)


class DriveAccessError(RuntimeError):
    """Raised when the service account cannot see the requested file/folder."""


class DriveClient:
    def __init__(self, service_account_file: str | None = None):
        # GOOGLE_SERVICE_ACCOUNT_JSON (inline key content) takes priority - for
        # platforms like Railway with no committed/persistent filesystem, set
        # this as a secret env var instead of shipping the key file. Falls
        # back to a local file path for local dev.
        inline_json = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON")
        if inline_json:
            import json

            info = json.loads(inline_json)
            creds = service_account.Credentials.from_service_account_info(info, scopes=SCOPES)
        else:
            path = service_account_file or os.environ.get("GOOGLE_SERVICE_ACCOUNT_FILE")
            if not path:
                raise RuntimeError(
                    "Neither GOOGLE_SERVICE_ACCOUNT_JSON nor GOOGLE_SERVICE_ACCOUNT_FILE is set "
                    "(check your .env file)"
                )
            creds = service_account.Credentials.from_service_account_file(path, scopes=SCOPES)
        self.service = build("drive", "v3", credentials=creds, cache_discovery=False)
        self.sheets = build("sheets", "v4", credentials=creds, cache_discovery=False)
        self._folder_cache: dict[str, dict] = {}  # folder id -> {"name":..., "parents":[...]}

    def search(self, query: str, page_size: int = 20) -> list[dict]:
        safe_query = query.replace("'", "\\'")
        q = f"(name contains '{safe_query}' or fullText contains '{safe_query}') and trashed = false"
        result = self.service.files().list(
            q=q,
            pageSize=page_size,
            fields="files(id,name,mimeType,modifiedTime,parents,webViewLink)",
        ).execute()
        return result.get("files", [])

    def list_all_files(self) -> list[dict]:
        """One paginated sweep of every non-trashed file/folder visible to this
        service account. Much faster than walking the folder tree one API call
        per folder (that approach took 30s+ per team, minutes overall) since
        it's ~1 call per 1000 items instead of ~1 call per folder.
        """
        files: list[dict] = []
        token = None
        while True:
            result = self.service.files().list(
                q="trashed = false",
                pageSize=1000,
                pageToken=token,
                fields="nextPageToken, files(id,name,mimeType,modifiedTime,parents,webViewLink)",
            ).execute()
            files.extend(result.get("files", []))
            token = result.get("nextPageToken")
            if not token:
                break
        return files

    def resolve_file_teams(self, all_files: list[dict] | None = None) -> tuple[dict[str, str], dict[str, dict]]:
        """Resolve each non-folder file to its TEAM_FOLDERS team, entirely from
        one bulk listing (no per-folder API calls). Returns (file_id -> team,
        file_id -> file dict). A file under a stale-marked ancestor folder
        (e.g. "회장단-old") is excluded.
        """
        all_files = all_files if all_files is not None else self.list_all_files()
        by_id = {f["id"]: f for f in all_files}
        team_by_root = {v: k for k, v in TEAM_FOLDERS.items()}

        resolved: dict[str, str | None] = {}

        def resolve_folder(folder_id: str, visiting: set) -> str | None:
            if folder_id in resolved:
                return resolved[folder_id]
            if folder_id in team_by_root:
                resolved[folder_id] = team_by_root[folder_id]
                return resolved[folder_id]
            if folder_id in visiting:
                return None  # cycle guard
            info = by_id.get(folder_id)
            if not info or _is_stale_folder(info.get("name", "")):
                resolved[folder_id] = None
                return None
            parents = info.get("parents") or []
            if not parents:
                resolved[folder_id] = None
                return None
            visiting.add(folder_id)
            team = resolve_folder(parents[0], visiting)
            resolved[folder_id] = team
            return team

        file_team: dict[str, str] = {}
        for f in all_files:
            if f["mimeType"] not in SUPPORTED_MIME_TYPES:
                continue  # skip images/zips/csvs/etc. before any folder-walk work
            parents = f.get("parents") or []
            if not parents:
                continue
            team = resolve_folder(parents[0], set())
            if team:
                file_team[f["id"]] = team
        return file_team, by_id

    def get_metadata(self, file_id: str) -> dict:
        return self.service.files().get(
            fileId=file_id, fields="id,name,mimeType,modifiedTime,webViewLink"
        ).execute()

    def get_path(self, parents: list[str] | None, max_depth: int = 6) -> str:
        """Breadcrumb of ancestor folder names, e.g. "회장단 > 24 총회 > 24.16 임시총회".

        Walks parents on demand (a few extra API calls per search result) since
        building this for the whole Drive up front is the 87s bulk listing -
        too slow for a single interactive search_drive call. Folder names/parents
        are cached on this client instance so repeat lookups are free.
        """
        names: list[str] = []
        current = (parents or [None])[0]
        depth = 0
        while current and depth < max_depth:
            if current not in self._folder_cache:
                try:
                    info = self.service.files().get(fileId=current, fields="name,parents").execute()
                    self._folder_cache[current] = {"name": info.get("name"), "parents": info.get("parents") or []}
                except Exception:
                    break
            entry = self._folder_cache[current]
            names.append(entry["name"])
            current = (entry["parents"] or [None])[0]
            depth += 1
        return " > ".join(reversed(names))

    def guess_generation(self, *texts: str) -> str | None:
        """Best-effort 기수 (generation number) extracted from a filename/path,
        e.g. "13기 임시총회.pdf" -> "13". Returns None if no pattern matches -
        callers should fall back to modifiedTime for recency in that case.
        """
        for text in texts:
            if not text:
                continue
            m = GENERATION_PATTERN.search(text)
            if m:
                return m.group(1)
        return None

    def read_spreadsheet(self, file_id: str, sheet_name: str = "") -> str:
        """Live read of a Google Sheet - always current, never a stale embedded
        snapshot. Without sheet_name, returns the list of tab names so the
        caller can pick one (budget/attendance sheets often have per-기수 tabs).
        """
        meta = self.sheets.spreadsheets().get(spreadsheetId=file_id).execute()
        tabs = [s["properties"]["title"] for s in meta.get("sheets", [])]
        if not sheet_name:
            return "탭 목록: " + ", ".join(tabs) + "\n(sheet_name을 지정해서 다시 호출하세요)"
        if sheet_name not in tabs:
            return f"'{sheet_name}' 탭을 찾을 수 없습니다. 탭 목록: {', '.join(tabs)}"
        values = self.sheets.spreadsheets().values().get(
            spreadsheetId=file_id, range=sheet_name
        ).execute().get("values", [])
        return "\n".join("\t".join(str(cell) for cell in row) for row in values)

    def read_file_text(self, file_id: str, mime_type: str | None = None) -> str:
        if mime_type is None:
            mime_type = self.get_metadata(file_id)["mimeType"]

        if mime_type in EXPORT_TEXT_MIME_TYPES:
            data = self.service.files().export(
                fileId=file_id, mimeType=EXPORT_TEXT_MIME_TYPES[mime_type]
            ).execute()
            return data.decode("utf-8") if isinstance(data, bytes) else data
        if mime_type == "application/pdf":
            return self._extract_pdf(file_id)
        if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return self._extract_docx(file_id)
        if mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
            return self._extract_pptx(file_id)
        return ""  # unsupported type (images, spreadsheets, shortcuts, etc.) - skip silently

    def _download(self, file_id: str) -> io.BytesIO:
        request = self.service.files().get_media(fileId=file_id)
        buf = io.BytesIO()
        downloader = MediaIoBaseDownload(buf, request)
        done = False
        while not done:
            _, done = downloader.next_chunk()
        buf.seek(0)
        return buf

    def _extract_pdf(self, file_id: str) -> str:
        import pdfplumber

        with pdfplumber.open(self._download(file_id)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)

    def _extract_docx(self, file_id: str) -> str:
        import docx

        d = docx.Document(self._download(file_id))
        return "\n".join(p.text for p in d.paragraphs)

    def _extract_pptx(self, file_id: str) -> str:
        from pptx import Presentation

        prs = Presentation(self._download(file_id))
        lines = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    lines.append(shape.text_frame.text)
        return "\n".join(lines)
